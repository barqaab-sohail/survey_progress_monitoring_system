from pathlib import Path
from docx import Document
import fitz
from pypdf import PdfReader
word=Path('output/manual/HAZECO_Field_Survey_User_Manual.docx')
pdf=Path('output/pdf/HAZECO_Field_Survey_User_Manual.pdf')
d=Document(word)
text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
required=['Install and open the app','Sign in and download assignments','Choose a team and feeder','Create the transformer survey','S/E observations and survey data','GPS coordinates and waypoint IDs','Save a draft and continue later','Review, validate, and submit','After submission: same feeder or Home','Work without internet','Sync and confirm every upload','Common problems and fixes','Meter Number','Consumer reference number','Take photo','Synced']
missing=[x for x in required if x.lower() not in text.lower()]
r=PdfReader(str(pdf)); info=fitz.open(pdf)
print('docx_bytes',word.stat().st_size,'pdf_bytes',pdf.stat().st_size,'pages',len(info),'pdf_text_pages',len(r.pages),'missing',missing)
print('page_numbers',[(i+1,('Page '+str(i+1)) in info[i].get_text()) for i in range(len(info))])
print('annotated_image',Path('docs/mobile/manual/screenshots/01-login-annotated.png').stat().st_size)
if missing or len(info)!=9 or any(('Page '+str(i+1)) not in info[i].get_text() for i in range(len(info))): raise SystemExit(1)
