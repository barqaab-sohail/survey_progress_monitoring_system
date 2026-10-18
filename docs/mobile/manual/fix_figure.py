from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
p='output/manual/HAZECO_Field_Survey_User_Manual.docx'; d=Document(p)
for para in d.paragraphs:
 if para._p.xpath('.//w:drawing'):
  para.paragraph_format.keep_together=False; para.paragraph_format.keep_with_next=False
  for ext in para._p.xpath('.//wp:extent'):
   ext.set('cx',str(int(int(ext.get('cx'))*.70))); ext.set('cy',str(int(int(ext.get('cy'))*.70)))
 if para.text.startswith('Figure 1.'):
  para.paragraph_format.keep_together=False; para.paragraph_format.keep_with_next=False
for table in d.tables:
 for row in table.rows:
  trPr=row._tr.get_or_add_trPr()
  if not trPr.xpath('./w:cantSplit'):trPr.append(OxmlElement('w:cantSplit'))
d.save(p); print('saved figure layout')
