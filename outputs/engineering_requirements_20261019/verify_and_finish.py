"""Add print-only XLSX metadata, then independently read and validate the deliverable.
Artifact Tool authored the workbook. XML is used only for print features absent from
the documented authoring surface; openpyxl is used read-only for verification.
"""
from pathlib import Path
import json, hashlib, zipfile, math
import xml.etree.ElementTree as ET
from pyproj import Transformer
from openpyxl import load_workbook

OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[1]
data=json.loads((OUT/'requirements_data.json').read_text(encoding='utf-8'))
layout=json.loads((OUT/'workbook_layout.json').read_text())
target=OUT/'Engineering_Configuration_Data_Requirement_Sheet.xlsx'
ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
ET.register_namespace('',ns)
def q(s):return '{'+ns+'}'+s

# Set landscape A3, a one-page print width, and repeated title/header rows.
with zipfile.ZipFile(target) as archive:
    parts={name:archive.read(name) for name in archive.namelist()}
book=ET.fromstring(parts['xl/workbook.xml'])
defined=book.find(q('definedNames'))
if defined is None:
    defined=ET.Element(q('definedNames'))
    position=next((i for i,c in enumerate(book) if c.tag==q('calcPr')),len(book))
    book.insert(position,defined)
for item in list(defined):
    if item.get('name') in ('_xlnm.Print_Area','_xlnm.Print_Titles'):
        defined.remove(item)
for idx,(name,info) in enumerate(layout.items()):
    file=f'xl/worksheets/sheet{idx+1}.xml'
    sheet=ET.fromstring(parts[file])
    props=sheet.find(q('sheetPr'))
    if props is None:
        props=ET.Element(q('sheetPr'));sheet.insert(0,props)
    setup_props=props.find(q('pageSetUpPr'))
    if setup_props is None:setup_props=ET.SubElement(props,q('pageSetUpPr'))
    setup_props.set('fitToPage','1')
    for name_to_remove in ['pageMargins','pageSetup','headerFooter']:
        element=sheet.find(q(name_to_remove))
        if element is not None:sheet.remove(element)
    ET.SubElement(sheet,q('pageMargins'),left='0.25',right='0.25',top='0.35',bottom='0.35',header='0.12',footer='0.12')
    ET.SubElement(sheet,q('pageSetup'),paperSize='8',orientation='landscape',fitToWidth='1',fitToHeight='0')
    footer=ET.SubElement(sheet,q('headerFooter'))
    ET.SubElement(footer,q('oddFooter')).text='&LHAZECO Engineering Requirements&CPage &P of &N&R&A'
    escaped=name.replace("'","''")
    for suffix,text in [('Print_Area',f"'{escaped}'!$A$1:${info['lastCol']}${info['lastRow']}"),('Print_Titles',f"'{escaped}'!$1:$5")]:
        ET.SubElement(defined,q('definedName'),name='_xlnm.'+suffix,localSheetId=str(idx)).text=text
    parts[file]=ET.tostring(sheet,encoding='utf-8',xml_declaration=True)
parts['xl/workbook.xml']=ET.tostring(book,encoding='utf-8',xml_declaration=True)
temp=target.with_suffix('.printing.xlsx')
with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as archive:
    for name,content in parts.items():archive.writestr(name,content)
temp.replace(target)

wb=load_workbook(target,data_only=False)
assert wb.sheetnames==list(layout)
seen=[]
for name,info in layout.items():
    sheet=wb[name]
    if info['inputRows']:
        assert [sheet.cell(5,c).value for c in range(1,9)]==data['headers']
        assert sheet.freeze_panes=='C6', (name,sheet.freeze_panes)
        selected=[r for r in data['rows'] if r['level']==name]
        for row,expected in zip(info['inputRows'],selected):
            values=[sheet.cell(row,c).value for c in range(1,9)]
            assert values==[expected['sr'],expected['name'],expected['purpose'],expected['format'],expected['sample'],expected['owner'],None,expected['remark']],(name,row)
            assert sheet.cell(row,7).data_type!='f'
            seen.append(expected['sr'])
        # Every cell in the response column, except its header, is truly blank.
        assert all(sheet.cell(r,7).value is None for r in range(1,info['lastRow']+1) if r!=5)
    assert sheet.page_setup.orientation=='landscape'
    assert int(sheet.page_setup.paperSize)==8
    assert sheet.page_setup.fitToWidth==1
    assert sheet.print_title_rows=='$1:$5'
    assert not sheet.sheet_view.showGridLines
    for row in sheet.iter_rows():
        for cell in row:
            assert cell.data_type!='e',(name,cell.coordinate,cell.value)
assert sorted(seen)==list(range(1,117))
assert len(seen)==len(set(seen))
with zipfile.ZipFile(target) as archive:
    workbook_xml=ET.fromstring(archive.read('xl/workbook.xml'))
    names=workbook_xml.find(q('definedNames'))
    assert len(names)==8
    assert len({(n.get('name'),n.get('localSheetId')) for n in names})==8

# Verify read-only source integrity and record the concrete projection conflict.
manifest=json.loads((ROOT/'resources/mdb/engineering-defaults.json').read_text())
source_root=Path('C:/Users/sohai/OneDrive/Desktop/HAZECO/Sample data')
for source in manifest['evidence']['sources']:
    assert hashlib.sha256((source_root/source['file']).read_bytes()).hexdigest()==source['sha256']
for source in json.loads((OUT/'reviewed_source_inventory.json').read_text()):
    path=source_root/source['file'] if 'file' in source else ROOT/source['path']
    assert hashlib.sha256(path.read_bytes()).hexdigest()==source['sha256']
for source in data['code_sources']:
    assert hashlib.sha256((ROOT/source['path']).read_bytes()).hexdigest()==source['sha256']
point=next(p for p in ET.parse(source_root/'A1408.gpx').getroot().iter() if p.tag.rsplit('}',1)[-1]=='wpt' and any(c.tag.rsplit('}',1)[-1]=='name' and c.text=='591' for c in p))
doc=json.loads((ROOT/'output/engineering-defaults-review/IndustrialT-4060116233.inspection.json').read_text(encoding='utf-8-sig'))
node=next(n for t in doc['tables'] if t['name']=='Node' for n in t['rows'] if n['NodeId']=='1140825591')
geo={'gpx':'A1408.gpx','waypoint':'591','latitude':float(point.attrib['lat']),'longitude':float(point.attrib['lon']),'mdb':'IndustrialT-4060116233.mdb','node':node,'comparisons':[]}
for epsg in [32642,32643]:
    x,y=Transformer.from_crs(4326,epsg,always_xy=True).transform(geo['longitude'],geo['latitude'])
    geo['comparisons'].append({'epsg':epsg,'x':x,'y':y,'distance_to_mdb_m':math.hypot(x-node['X'],y-node['Y'])})
(OUT/'geo_check.json').write_text(json.dumps(geo,indent=2))
result={'file':str(target),'requirements':116,'mandatory':33,'conditional':62,'optional':21,'actual_value_cells_all_blank':True,'column_order_verified':True,'unique_serial_numbers_verified':True,'error_cells':0,'freeze_panes':'C6','print_format':'A3 landscape, one page wide, repeating title/header rows','application_code_hashes_unchanged':True,'reference_and_survey_source_hashes_unchanged':True,'bytes':target.stat().st_size}
(OUT/'verification.json').write_text(json.dumps(result,indent=2))
print(json.dumps(result,indent=2))
