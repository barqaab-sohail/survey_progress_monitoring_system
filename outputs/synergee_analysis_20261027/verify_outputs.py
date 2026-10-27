import json, pathlib, math, hashlib, collections
import openpyxl
ROOT=pathlib.Path(__file__).resolve().parent
p=json.loads((ROOT/'analysis_data.json').read_text(encoding='utf-8'))
book=ROOT/'Synergee_MDB_Entry_Reference.xlsx'
w=openpyxl.load_workbook(book,read_only=True,data_only=False)
errors=[]
def eq(a,b):
    if isinstance(a,(int,float)) and isinstance(b,(int,float)):
        return math.isclose(a,b,rel_tol=1e-14,abs_tol=1e-12)
    return type(a)==type(b) and a==b
for i,(got,expected) in enumerate(zip(w['All entries'].iter_rows(min_row=6,max_row=len(p['entries'])+5,values_only=True),p['entries']),6):
    for j,(a,b) in enumerate(zip(got,expected),1):
        if j!=6 and a is None and b=='':continue
        if not eq(a,b):errors.append(('All entries',i,j,a,b))
assert w['All entries']['A13185'].value==p['entries'][-1][0]
assert w['All entries']['A13186'].value is None
assert w['Field dictionary']['A262'].value==p['fields'][-1][0]
for name,key in [('Settings','settings'),('Sections','sections'),('Nodes','nodes')]:
    rows=p[key]
    for i,(got,expected) in enumerate(zip(w[name].iter_rows(min_row=6,max_row=len(rows)+5,values_only=True),rows),6):
        for j,(a,b) in enumerate(zip(got,expected),1):
            # Excel serializes empty strings as blank cells in this curated view.
            if a is None and b=='':continue
            if not eq(a,b):errors.append((name,i,j,a,b))
for d in p['summaries']:
    source=pathlib.Path('C:/Users/sohai/OneDrive/Desktop/HAZECO')/d['file']
    assert hashlib.sha256(source.read_bytes()).hexdigest().upper()==d['hash']
for row in w['All entries'].iter_rows(min_row=6):
    assert row[5].data_type!='f',('Raw value became formula',row[5].coordinate)
assert not errors,errors[:12]
wc=openpyxl.load_workbook(book,read_only=True,data_only=True)
for i,s in enumerate(p['summaries'],6):
    for col,key in [('G','customers'),('H','kva'),('I','length'),('K','dg_kw')]:
        assert eq(wc['Overview'][f'{col}{i}'].value,s[key]),(col,i)
assert w['Loads']['B6'].data_type=='s'
assert w['Loads']['B6'].value=='1059043030216724.2'
assert w['Nodes']['B8'].value=='3030216724.1'
assert w['Nodes']['B9'].value=='3030216724.2'
print('PASS: all 13,180 entry values and associated explanations match the extraction; curated settings, sections and nodes reconcile; identifiers remain exact text; summary formula caches reconcile; all seven source hashes remain unchanged.')
